"""
Mizan.ai — VAT Compliance Center report generation: job 1 (invoice
compliance check) and job 2 (period VAT return preparation). Sibling to
features/calculator/report.py rather than an addition to it: that module's
functions render an already-finalized single-invoice calculation result;
these render a batch of independently-checked invoices or a period's
aggregated figures, different shapes entirely. Follows the same
conventions though — per-format functions, static bilingual labels,
free-text translated via features/common/translation_client.maybe_translate,
never failing the whole report over one untranslated string.
"""

import io
from typing import Any, Dict, List, Literal, Tuple

ReportFormat = Literal["pdf", "docx", "xlsx"]
ReportLanguage = Literal["ar", "en"]

_LABELS: Dict[str, Dict[str, str]] = {
    "invoice_compliance_summary": {"en": "Invoice Compliance Summary", "ar": "ملخص الامتثال للفواتير"},
    "filename": {"en": "File", "ar": "الملف"},
    "status": {"en": "Status", "ar": "الحالة"},
    "reason": {"en": "Reason", "ar": "السبب"},
    "clean": {"en": "Clean", "ar": "سليم"},
    "flagged": {"en": "Needs review", "ar": "يحتاج مراجعة"},
    "totals": {"en": "Total: {total} — Clean: {clean} — Flagged: {flagged}",
               "ar": "الإجمالي: {total} — سليم: {clean} — يحتاج مراجعة: {flagged}"},
    "invoice_issues": {"en": "Invoice Issues", "ar": "مشاكل الفاتورة"},
    "extracted_fields": {"en": "Extracted fields", "ar": "الحقول المستخرجة"},
    "structural_issues": {"en": "Structural issues", "ar": "مشاكل هيكلية"},
    "numeric_mismatches": {"en": "Numeric mismatches", "ar": "فروقات في الأرقام"},
    "explanation": {"en": "Explanation", "ar": "التفسير"},
    "document_value": {"en": "Document shows", "ar": "تظهر الوثيقة"},
    "recalculated_value": {"en": "Recalculated value", "ar": "القيمة المعاد حسابها"},
    "rounding_difference_only": {"en": "Rounding difference only — not treated as an issue", "ar": "فرق تقريب فقط — لا يُعتبر مشكلة"},
    "seller_name": {"en": "Seller name", "ar": "اسم البائع"},
    "seller_vat_number": {"en": "Seller VAT number", "ar": "الرقم الضريبي للبائع"},
    "buyer_name": {"en": "Buyer name", "ar": "اسم المشتري"},
    "issue_date": {"en": "Issue date", "ar": "تاريخ الإصدار"},
    "tax_exclusive_amount": {"en": "Amount before VAT", "ar": "المبلغ قبل الضريبة"},
    "tax_amount": {"en": "VAT amount", "ar": "مبلغ الضريبة"},
    "tax_inclusive_amount": {"en": "Total (incl. VAT)", "ar": "الإجمالي شامل الضريبة"},
    # Job 2 — period VAT return preparation
    "vat_return_filing": {"en": "VAT Return — Filing Figures", "ar": "إقرار ضريبة القيمة المضافة — أرقام الإقرار"},
    "vat_return_workpapers": {"en": "VAT Return — Workpapers", "ar": "إقرار ضريبة القيمة المضافة — أوراق العمل"},
    "output_vat": {"en": "Output VAT (on sales)", "ar": "ضريبة المخرجات (على المبيعات)"},
    "input_vat_total": {"en": "Input VAT, all purchases", "ar": "ضريبة المدخلات، جميع المشتريات"},
    "reclaimable_input_vat": {"en": "Reclaimable Input VAT", "ar": "ضريبة المدخلات القابلة للاسترداد"},
    "net_position": {"en": "Net VAT position (owed if positive, refund if negative)",
                      "ar": "صافي مركز ضريبة القيمة المضافة (مستحق إذا كان موجباً، مسترد إذا كان سالباً)"},
    "excluded_purchases": {"en": "Purchases excluded from reclaimable VAT", "ar": "المشتريات المستبعدة من ضريبة المدخلات القابلة للاسترداد"},
    "source": {"en": "Source", "ar": "المصدر"},
    "description": {"en": "Description", "ar": "الوصف"},
    "vat_amount_col": {"en": "VAT amount", "ar": "مبلغ الضريبة"},
    "reason_col": {"en": "Reason", "ar": "السبب"},
    "consistency_issues": {"en": "Consistency checks", "ar": "فحوصات الاتساق"},
}


def _label(key: str, language: ReportLanguage) -> str:
    return _LABELS.get(key, {}).get(language, key)


def _maybe_translate(text: str, target_language: ReportLanguage) -> str:
    from features.calculator.config import TRANSLATOR_INTERNAL_URL
    from features.common.translation_client import maybe_translate

    return maybe_translate(text, target_language, TRANSLATOR_INTERNAL_URL)


def _flagged_reason_summary(invoice: Dict[str, Any]) -> str:
    """ALL error-severity structural issues, not just the first one — a
    real invoice can be flagged for more than one reason at once (e.g. a
    malformed seller VAT number AND a malformed buyer VAT number
    simultaneously), and a reader glancing at just this one column should
    see every reason, not silently lose everything after the first. Found
    via live testing: the previous first-issue-only version made a real,
    still-present second issue look like it had been dropped entirely,
    even though the full per-invoice detail report always had it."""
    messages = [i["message"] for i in invoice.get("structural_issues", []) if i.get("severity") == "error"]
    significant_mismatches = [m for m in invoice.get("mismatches", []) if not m.get("within_tolerance")]
    if significant_mismatches:
        messages.append("Recalculated VAT/total doesn't match what's printed on the document.")
    return " | ".join(messages)


def _summary_reason(invoice: Dict[str, Any], language: ReportLanguage) -> str:
    """A flagged invoice gets its full reason summary. A CLEAN invoice
    normally gets nothing (there's nothing to say) — except when it has a
    within-tolerance mismatch, which stays soft-labeled rather than fully
    silent: a real difference was found and deliberately excused, and that
    should stay visible somewhere even though it wasn't treated as a
    problem, consistent with this feature never silently deciding
    something is fine without showing its reasoning."""
    if invoice["status"] == "flagged":
        return _flagged_reason_summary(invoice)
    within_tolerance = [m for m in invoice.get("mismatches", []) if m.get("within_tolerance")]
    if within_tolerance:
        fields = ", ".join(m["field"] for m in within_tolerance)
        return f"{_label('rounding_difference_only', language)} ({fields})"
    return ""


# ---------------------------------------------------------------------------
# Summary report — every invoice, one row each
# ---------------------------------------------------------------------------


def _summary_xlsx(invoices: List[Dict[str, Any]], language: ReportLanguage) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _label("invoice_compliance_summary", language)[:31]
    ws.append([_label("filename", language), _label("status", language), _label("reason", language)])
    for inv in invoices:
        status_label = _label(inv["status"], language)
        reason = _maybe_translate(_summary_reason(inv, language), language)
        ws.append([inv["filename"], status_label, reason])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _summary_docx(invoices: List[Dict[str, Any]], language: ReportLanguage) -> bytes:
    import docx

    doc = docx.Document()
    doc.add_heading(_label("invoice_compliance_summary", language), level=1)
    table = doc.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text = _label("filename", language)
    header[1].text = _label("status", language)
    header[2].text = _label("reason", language)
    for inv in invoices:
        row = table.add_row().cells
        row[0].text = inv["filename"]
        row[1].text = _label(inv["status"], language)
        row[2].text = _maybe_translate(_summary_reason(inv, language), language)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _summary_pdf(invoices: List[Dict[str, Any]], language: ReportLanguage) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8.5, leading=11)
    story = [Paragraph(_label("invoice_compliance_summary", language), styles["Title"]), Spacer(1, 12)]

    # Free-text cells (filename, reason) are wrapped in Paragraph objects,
    # not plain strings — a plain string in a reportlab Table cell doesn't
    # wrap at all, it just overflows past the cell/column boundary, which
    # is exactly what produced the "columns merging" symptom reported live
    # for a long reason sentence or filename.
    data = [[_label("filename", language), _label("status", language), _label("reason", language)]]
    for inv in invoices:
        reason = _maybe_translate(_summary_reason(inv, language), language)
        data.append([Paragraph(inv["filename"], cell_style), _label(inv["status"], language), Paragraph(reason, cell_style)])

    # Column widths total 15cm — A4 is 21cm wide; SimpleDocTemplate's
    # default 1-inch (2.54cm) margins on each side leave ~15.9cm usable,
    # so this fits with a safety margin (the previous 17cm total didn't,
    # which was the actual root cause — reportlab draws a Table at exactly
    # the widths given, it does not shrink to fit the page).
    table = Table(data, colWidths=[5 * cm, 2.5 * cm, 7.5 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16213e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()


def generate_invoice_check_summary(invoices: List[Dict[str, Any]], format: ReportFormat = "xlsx", language: ReportLanguage = "en") -> bytes:
    if format == "xlsx":
        return _summary_xlsx(invoices, language)
    if format == "docx":
        return _summary_docx(invoices, language)
    if format == "pdf":
        return _summary_pdf(invoices, language)
    raise ValueError(f"Unsupported report format: {format!r}")


# ---------------------------------------------------------------------------
# Per-invoice detail — only ever generated for a flagged invoice
# ---------------------------------------------------------------------------


def generate_invoice_issue_detail(invoice: Dict[str, Any], format: ReportFormat = "pdf", language: ReportLanguage = "en") -> bytes:
    """Always PDF-shaped as a narrative regardless of `format` — this is a
    single flagged invoice's explanation, not a data table (same reasoning
    as report.py's issues_summary report_type)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph(invoice["filename"], styles["Title"]), Spacer(1, 12)]

    extracted = invoice.get("extracted") or {}
    if extracted:
        story.append(Paragraph(_label("extracted_fields", language), styles["Heading2"]))
        for key in ("seller_name", "seller_vat_number", "buyer_name", "issue_date",
                    "tax_exclusive_amount", "tax_amount", "tax_inclusive_amount"):
            value = extracted.get(key)
            if value is not None:
                story.append(Paragraph(f"{_label(key, language)}: {value}", styles["Normal"]))
        story.append(Spacer(1, 12))

    if invoice.get("structural_issues"):
        story.append(Paragraph(_label("structural_issues", language), styles["Heading2"]))
        for issue in invoice["structural_issues"]:
            message = _maybe_translate(issue.get("message", ""), language)
            story.append(Paragraph(f"[{issue.get('severity', 'error')}] {message}", styles["Normal"]))
        story.append(Spacer(1, 12))

    if invoice.get("mismatches"):
        story.append(Paragraph(_label("numeric_mismatches", language), styles["Heading2"]))
        for m in invoice["mismatches"]:
            if m.get("within_tolerance"):
                # Shown, not hidden — but clearly labeled as harmless
                # rounding noise rather than presented identically to a
                # real discrepancy (see MISMATCH_TOLERANCE in config.py).
                story.append(Paragraph(
                    f"{m.get('field')}: {_label('rounding_difference_only', language)} "
                    f"({_label('document_value', language)} {m.get('document_value')}, "
                    f"{_label('recalculated_value', language)} {m.get('recalculated_value')})",
                    styles["Normal"],
                ))
            else:
                story.append(Paragraph(
                    f"{m.get('field')}: {_label('document_value', language)} {m.get('document_value')}, "
                    f"{_label('recalculated_value', language)} {m.get('recalculated_value')}",
                    styles["Normal"],
                ))
        story.append(Spacer(1, 12))

    if invoice.get("explanation"):
        story.append(Paragraph(_label("explanation", language), styles["Heading2"]))
        story.append(Paragraph(_maybe_translate(invoice["explanation"], language), styles["Normal"]))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Job 2 — period VAT return preparation. Two reports from the same result
# dict (see period_return.py's process_period_return()/
# apply_review_decisions()): "filing" is just the headline figures, shaped
# to transcribe into an actual return; "workpapers" is the full trail —
# same headline plus every excluded purchase with its reason and every
# consistency-check finding, so an auditor can see how the number was
# built, not just trust it.
# ---------------------------------------------------------------------------


def _headline_rows(result: Dict[str, Any], language: ReportLanguage) -> List[Tuple[str, str]]:
    return [
        (_label("output_vat", language), result["output_vat"]),
        (_label("input_vat_total", language), result["input_vat_total"]),
        (_label("reclaimable_input_vat", language), result["reclaimable_input_vat"]),
        (_label("net_position", language), result["net_position"]),
    ]


def _filing_xlsx(result: Dict[str, Any], period_label: str, language: ReportLanguage) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _label("vat_return_filing", language)[:31]
    ws.append([_label("vat_return_filing", language), period_label])
    ws.append([])
    for label, value in _headline_rows(result, language):
        ws.append([label, value])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _filing_docx(result: Dict[str, Any], period_label: str, language: ReportLanguage) -> bytes:
    import docx

    doc = docx.Document()
    doc.add_heading(f"{_label('vat_return_filing', language)} — {period_label}", level=1)
    for label, value in _headline_rows(result, language):
        doc.add_paragraph(f"{label}: {value}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _filing_pdf(result: Dict[str, Any], period_label: str, language: ReportLanguage) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"{_label('vat_return_filing', language)} — {period_label}", styles["Title"]),
        Spacer(1, 12),
    ]
    for label, value in _headline_rows(result, language):
        story.append(Paragraph(f"{label}: {value}", styles["Normal"]))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()


def generate_period_return_filing(
    result: Dict[str, Any], period_label: str, format: ReportFormat = "xlsx", language: ReportLanguage = "en",
) -> bytes:
    if format == "xlsx":
        return _filing_xlsx(result, period_label, language)
    if format == "docx":
        return _filing_docx(result, period_label, language)
    if format == "pdf":
        return _filing_pdf(result, period_label, language)
    raise ValueError(f"Unsupported report format: {format!r}")


def _excluded_purchases(result: Dict[str, Any]) -> List[Dict[str, Any]]:
    return [row for row in result.get("purchase_rows", []) if not row["reclaimable"]]


def _workpapers_xlsx(result: Dict[str, Any], period_label: str, language: ReportLanguage) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _label("vat_return_workpapers", language)[:31]
    ws.append([_label("vat_return_workpapers", language), period_label])
    ws.append([])
    for label, value in _headline_rows(result, language):
        ws.append([label, value])

    ws.append([])
    ws.append([_label("excluded_purchases", language)])
    ws.append([_label("source", language), _label("description", language),
               _label("vat_amount_col", language), _label("reason_col", language)])
    for row in _excluded_purchases(result):
        ws.append([row["source"], _maybe_translate(row["description"], language), row["vat_amount"],
                   _maybe_translate(row["reason"], language)])

    if result.get("consistency_issues"):
        ws.append([])
        ws.append([_label("consistency_issues", language)])
        for issue in result["consistency_issues"]:
            ws.append([issue["source"], _maybe_translate(issue["message"], language)])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _workpapers_docx(result: Dict[str, Any], period_label: str, language: ReportLanguage) -> bytes:
    import docx

    doc = docx.Document()
    doc.add_heading(f"{_label('vat_return_workpapers', language)} — {period_label}", level=1)
    for label, value in _headline_rows(result, language):
        doc.add_paragraph(f"{label}: {value}")

    excluded = _excluded_purchases(result)
    if excluded:
        doc.add_heading(_label("excluded_purchases", language), level=2)
        table = doc.add_table(rows=1, cols=4)
        header = table.rows[0].cells
        header[0].text = _label("source", language)
        header[1].text = _label("description", language)
        header[2].text = _label("vat_amount_col", language)
        header[3].text = _label("reason_col", language)
        for row in excluded:
            cells = table.add_row().cells
            cells[0].text = row["source"]
            cells[1].text = _maybe_translate(row["description"], language)
            cells[2].text = row["vat_amount"]
            cells[3].text = _maybe_translate(row["reason"], language)

    if result.get("consistency_issues"):
        doc.add_heading(_label("consistency_issues", language), level=2)
        for issue in result["consistency_issues"]:
            doc.add_paragraph(f"{issue['source']}: {_maybe_translate(issue['message'], language)}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _workpapers_pdf(result: Dict[str, Any], period_label: str, language: ReportLanguage) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8.5, leading=11)
    story = [
        Paragraph(f"{_label('vat_return_workpapers', language)} — {period_label}", styles["Title"]),
        Spacer(1, 12),
    ]
    for label, value in _headline_rows(result, language):
        story.append(Paragraph(f"{label}: {value}", styles["Normal"]))
    story.append(Spacer(1, 12))

    excluded = _excluded_purchases(result)
    if excluded:
        story.append(Paragraph(_label("excluded_purchases", language), styles["Heading2"]))
        data = [[_label("source", language), _label("description", language),
                 _label("vat_amount_col", language), _label("reason_col", language)]]
        for row in excluded:
            data.append([
                Paragraph(row["source"], cell_style), Paragraph(_maybe_translate(row["description"], language), cell_style),
                row["vat_amount"], Paragraph(_maybe_translate(row["reason"], language), cell_style),
            ])
        # Column widths total 15cm — see _summary_pdf's comment on why this
        # matters: reportlab draws a Table at exactly the widths given, so
        # anything over the ~15.9cm usable width on A4 (with default
        # margins) overflows the page rather than shrinking to fit. Cells
        # are also Paragraph-wrapped, not plain strings, so long text wraps
        # within its column instead of bleeding into the next one.
        table = Table(data, colWidths=[3.5 * cm, 3 * cm, 2.5 * cm, 6 * cm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16213e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
        story.append(Spacer(1, 12))

    if result.get("consistency_issues"):
        story.append(Paragraph(_label("consistency_issues", language), styles["Heading2"]))
        for issue in result["consistency_issues"]:
            story.append(Paragraph(f"{issue['source']}: {_maybe_translate(issue['message'], language)}", styles["Normal"]))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()


def generate_period_return_workpapers(
    result: Dict[str, Any], period_label: str, format: ReportFormat = "xlsx", language: ReportLanguage = "en",
) -> bytes:
    if format == "xlsx":
        return _workpapers_xlsx(result, period_label, language)
    if format == "docx":
        return _workpapers_docx(result, period_label, language)
    if format == "pdf":
        return _workpapers_pdf(result, period_label, language)
    raise ValueError(f"Unsupported report format: {format!r}")

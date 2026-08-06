"""
Mizan.ai — report generation (Feature E, Step 7). Deterministic rendering,
not a judgment call — renders an already-finalized, already-confirmed
calculation result. Never anything model-derived (same "our data wins"
invariant as calculate_vat/validate_zatca_form).

Two distinct report types:
  - "data": the line-item/totals export. Default XLSX, PDF/DOCX optional.
  - "issues_summary": always PDF — a narrative explaining structural
    issues + numeric mismatches + the QwenBrain explanation, not a data
    table. Only meaningful when there's something to explain.

Numbers never change by language — only labels/descriptions are
translated, via an HTTP call to the Translator service's internal endpoint
(Calculator and Translator are separate services — this can no longer be
an in-process function call). If translation fails for any one
description, that description falls back to its original text rather than
failing the whole report — a report with one untranslated line is far
better than no report.

KNOWN RISK, not yet resolved: reportlab has no native right-to-left Arabic
text shaping. Arabic PDF output needs real testing against real Arabic
data before it can be trusted — this module does not claim to solve that,
it renders best-effort and callers should treat Arabic PDF output as
unverified until tested.
"""

import io
from typing import Any, Dict, List, Literal, Optional

from .engine import CalculationResult

ReportFormat = Literal["pdf", "docx", "xlsx"]
ReportLanguage = Literal["ar", "en"]
ReportType = Literal["data", "issues_summary"]

# Bilingual label templates for standard report fields — static, not
# translated on the fly (numbers/labels are structural, not free text).
_LABELS: Dict[str, Dict[str, str]] = {
    "invoice_total": {"en": "Invoice Total", "ar": "إجمالي الفاتورة"},
    "subtotal": {"en": "Subtotal", "ar": "المجموع الفرعي"},
    "total_vat": {"en": "Total VAT", "ar": "إجمالي ضريبة القيمة المضافة"},
    "description": {"en": "Description", "ar": "الوصف"},
    "taxable_base": {"en": "Taxable Amount", "ar": "المبلغ الخاضع للضريبة"},
    "tax_category": {"en": "Tax Category", "ar": "فئة الضريبة"},
    "vat_amount": {"en": "VAT Amount", "ar": "مبلغ ضريبة القيمة المضافة"},
    "structural_issues": {"en": "Structural Issues", "ar": "مشاكل هيكلية"},
    "numeric_mismatches": {"en": "Numeric Mismatches", "ar": "فروقات في الأرقام"},
    "explanation": {"en": "Explanation", "ar": "التفسير"},
    "document_value": {"en": "Document shows", "ar": "تظهر الوثيقة"},
    "recalculated_value": {"en": "Recalculated value", "ar": "القيمة المعاد حسابها"},
}


def _label(key: str, language: ReportLanguage) -> str:
    return _LABELS.get(key, {}).get(language, key)


def _maybe_translate(text: str, target_language: ReportLanguage) -> str:
    """Thin wrapper around the shared features/common/translation_client —
    also used by features/comparator/report.py, extracted here rather
    than duplicated once a second service needed the exact same logic."""
    from features.common.translation_client import maybe_translate

    from .config import TRANSLATOR_INTERNAL_URL

    return maybe_translate(text, target_language, TRANSLATOR_INTERNAL_URL)


# ---------------------------------------------------------------------------
# "data" report type
# ---------------------------------------------------------------------------


def _generate_data_xlsx(result: CalculationResult, language: ReportLanguage) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        _label("description", language), _label("taxable_base", language),
        _label("tax_category", language), _label("vat_amount", language),
    ])
    for line in result.lines:
        description = _maybe_translate(line.description, language)
        ws.append([description, float(line.taxable_base), line.tax_category.value, float(line.vat_amount)])

    ws.append([])
    ws.append([_label("subtotal", language), float(result.subtotal)])
    ws.append([_label("total_vat", language), float(result.total_vat)])
    ws.append([_label("invoice_total", language), float(result.grand_total)])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generate_data_docx(result: CalculationResult, language: ReportLanguage) -> bytes:
    import docx

    doc = docx.Document()
    doc.add_heading(_label("invoice_total", language), level=1)

    table = doc.add_table(rows=1, cols=4)
    header = table.rows[0].cells
    header[0].text = _label("description", language)
    header[1].text = _label("taxable_base", language)
    header[2].text = _label("tax_category", language)
    header[3].text = _label("vat_amount", language)

    for line in result.lines:
        row = table.add_row().cells
        row[0].text = _maybe_translate(line.description, language)
        row[1].text = str(line.taxable_base)
        row[2].text = line.tax_category.value
        row[3].text = str(line.vat_amount)

    doc.add_paragraph(f"{_label('subtotal', language)}: {result.subtotal}")
    doc.add_paragraph(f"{_label('total_vat', language)}: {result.total_vat}")
    doc.add_paragraph(f"{_label('invoice_total', language)}: {result.grand_total}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _generate_data_pdf(result: CalculationResult, language: ReportLanguage) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    story = [Paragraph(_label("invoice_total", language), styles["Title"]), Spacer(1, 12)]

    data = [[
        _label("description", language), _label("taxable_base", language),
        _label("tax_category", language), _label("vat_amount", language),
    ]]
    for line in result.lines:
        data.append([
            _maybe_translate(line.description, language), str(line.taxable_base),
            line.tax_category.value, str(line.vat_amount),
        ])

    table = Table(data, colWidths=[7 * cm, 3.5 * cm, 3.5 * cm, 3.5 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16213e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    story.append(table)
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"{_label('subtotal', language)}: {result.subtotal}", styles["Normal"]))
    story.append(Paragraph(f"{_label('total_vat', language)}: {result.total_vat}", styles["Normal"]))
    story.append(Paragraph(f"{_label('invoice_total', language)}: {result.grand_total}", styles["Normal"]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    doc.build(story)
    return buf.getvalue()


def generate_data_report(result: CalculationResult, format: ReportFormat = "xlsx", language: ReportLanguage = "en") -> bytes:
    if format == "xlsx":
        return _generate_data_xlsx(result, language)
    if format == "docx":
        return _generate_data_docx(result, language)
    if format == "pdf":
        return _generate_data_pdf(result, language)
    raise ValueError(f"Unsupported report format: {format!r}")


# ---------------------------------------------------------------------------
# "issues_summary" report type — always PDF, narrative
# ---------------------------------------------------------------------------


def generate_issues_summary(
    structural_issues: List[Dict[str, Any]],
    mismatches: List[Dict[str, Any]],
    explanation: Optional[str],
    language: ReportLanguage = "en",
) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    story = [Paragraph(_label("structural_issues", language), styles["Title"]), Spacer(1, 12)]

    if structural_issues:
        story.append(Paragraph(_label("structural_issues", language), styles["Heading2"]))
        for issue in structural_issues:
            story.append(Paragraph(f"[{issue.get('severity', 'error')}] {issue.get('message', '')}", styles["Normal"]))
        story.append(Spacer(1, 12))

    if mismatches:
        story.append(Paragraph(_label("numeric_mismatches", language), styles["Heading2"]))
        for m in mismatches:
            story.append(Paragraph(
                f"{m.get('field')}: {_label('document_value', language)} {m.get('document_value')}, "
                f"{_label('recalculated_value', language)} {m.get('recalculated_value')}",
                styles["Normal"],
            ))
        story.append(Spacer(1, 12))

    if explanation:
        story.append(Paragraph(_label("explanation", language), styles["Heading2"]))
        story.append(Paragraph(explanation, styles["Normal"]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    doc.build(story)
    return buf.getvalue()

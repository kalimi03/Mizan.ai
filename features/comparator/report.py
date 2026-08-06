"""
Mizan.ai — Comparator (Feature C) report generation (handoff doc Step 5).
Deterministic rendering of an already-finalized ReconciliationResult — same
"our data wins" invariant as features/calculator/report.py, and follows
its exact per-format-function shape (xlsx/docx/pdf), but is its own module
rather than importing that one: the underlying data shape (reconciliation
exceptions vs. VAT line items) is different, only the libraries and the
bilingual-label-template pattern are shared.

Numbers never change by language — only descriptions/labels are
translated, via features/common/translation_client.maybe_translate() (the
same shared helper features/calculator/report.py now also uses).

Styled to read like a real accounting reconciliation workpaper (dark
header banner, bold/shaded table headers, currency-formatted numbers,
status-colored exception rows) rather than a plain data dump — modeled on
the shape a finance team's own Excel-based month-end reconciliation
already takes: a summary/control-total section first (so a reviewer can
see at a glance how much of each file is accounted for), then the matched
detail, then the exceptions that need a decision.
"""

import io
from datetime import datetime
from decimal import Decimal
from typing import Dict, List, Literal, Optional, Tuple

from .config import TRANSLATOR_INTERNAL_URL
from .models import ReconciliationResult, ReviewItem

ReportFormat = Literal["pdf", "docx", "xlsx"]
ReportLanguage = Literal["ar", "en"]

# Same dark navy used by the pre-existing PDF header banner — kept as the
# one accent color across all three formats for a consistent look.
_ACCENT_HEX = "16213E"
_ACCENT_TEXT_HEX = "FFFFFF"
_CONFIRMED_HEX = "E2F0D9"    # light green
_DISMISSED_HEX = "F2DCDB"    # light red
_UNRESOLVED_HEX = "FFF2CC"   # light amber
_MISMATCH_HEX = "FCE4D6"     # light orange — draws the eye to the one thing that most needs it

_LABELS = {
    "reconciliation_report": {"en": "Reconciliation Report", "ar": "تقرير التسوية"},
    "generated": {"en": "Generated", "ar": "تاريخ الإصدار"},
    "summary": {"en": "Summary", "ar": "الملخص"},
    "sap_odoo_side": {"en": "SAP/Odoo Export", "ar": "تصدير ساب/أودو"},
    "second_doc_side": {"en": "Second Document", "ar": "المستند الثاني"},
    "grand_total": {"en": "Total in File", "ar": "إجمالي الملف"},
    "matched_amount": {"en": "Matched Amount", "ar": "المبلغ المطابق"},
    "matched_pct": {"en": "% Reconciled", "ar": "نسبة التسوية"},
    "total_matched": {"en": "Matched Items", "ar": "العناصر المطابقة"},
    "total_amount_mismatch": {"en": "Amount Mismatches", "ar": "حالات عدم تطابق المبلغ"},
    "total_group_match": {"en": "Group Matches (Consolidated)", "ar": "مطابقات مجمعة (توحيد فواتير)"},
    "total_sum_match": {"en": "Possible Groupings (Amount Only)", "ar": "تجميعات محتملة (بالمبلغ فقط)"},
    "total_other_exceptions": {"en": "Other Exceptions", "ar": "استثناءات أخرى"},
    "total_confirmed": {"en": "Confirmed", "ar": "مؤكد"},
    "total_dismissed": {"en": "Dismissed", "ar": "مرفوض"},
    "total_unresolved": {"en": "Unresolved", "ar": "غير محلول"},
    "matched_items": {"en": "Matched Items", "ar": "العناصر المتطابقة"},
    "exceptions": {"en": "Exceptions — Needs Review", "ar": "الاستثناءات - بحاجة إلى مراجعة"},
    "reference": {"en": "Reference", "ar": "المرجع"},
    "date": {"en": "Date", "ar": "التاريخ"},
    "amount": {"en": "Amount", "ar": "المبلغ"},
    "description": {"en": "Description", "ar": "الوصف"},
    "match_type": {"en": "Match Type", "ar": "نوع المطابقة"},
    "exception_type": {"en": "Exception Type", "ar": "نوع الاستثناء"},
    "sap_odoo_amount": {"en": "SAP/Odoo Amount", "ar": "مبلغ ساب/أودو"},
    "second_doc_amount": {"en": "Second Doc. Amount", "ar": "مبلغ المستند الثاني"},
    "difference": {"en": "Difference", "ar": "الفرق"},
    "status": {"en": "Status", "ar": "الحالة"},
    "explanation": {"en": "Explanation", "ar": "التفسير"},
    "confirmed": {"en": "Confirmed", "ar": "مؤكد"},
    "dismissed": {"en": "Dismissed", "ar": "مرفوض"},
    "unresolved": {"en": "Unresolved", "ar": "غير محلول"},
    "narration": {"en": "Summary Explanation", "ar": "ملخص التفسير"},
}

_MATCH_TYPE_LABELS = {
    "reference": {"en": "Reference match", "ar": "مطابقة بالمرجع"},
    "amount_unique": {"en": "Amount match", "ar": "مطابقة بالمبلغ"},
    "amount_date_fuzzy": {"en": "Amount + date match", "ar": "مطابقة بالمبلغ والتاريخ"},
}

_REVIEW_KIND_LABELS = {
    "amount_mismatch": {"en": "Amount mismatch", "ar": "عدم تطابق المبلغ"},
    "fuzzy_match": {"en": "Possible match", "ar": "تطابق محتمل"},
    "ambiguous": {"en": "Multiple candidates", "ar": "مرشحون متعددون"},
    "unmatched": {"en": "Unmatched", "ar": "غير مطابق"},
    "group_match": {"en": "Group match (consolidated)", "ar": "مطابقة مجمعة (توحيد فواتير)"},
    "sum_match": {"en": "Possible grouping — amount only, no other signal", "ar": "تجميع محتمل - بالمبلغ فقط، دون إشارة أخرى"},
}


def _label(key: str, language: ReportLanguage) -> str:
    return _LABELS.get(key, {}).get(language, key)


def _match_type_label(match_type: str, language: ReportLanguage) -> str:
    return _MATCH_TYPE_LABELS.get(match_type, {}).get(language, match_type)


def _review_kind_label(kind: str, language: ReportLanguage) -> str:
    return _REVIEW_KIND_LABELS.get(kind, {}).get(language, kind)


def _maybe_translate(text: Optional[str], target_language: ReportLanguage) -> str:
    if not text:
        return ""
    from features.common.translation_client import maybe_translate

    return maybe_translate(text, target_language, TRANSLATOR_INTERNAL_URL)


def _status_label(item: ReviewItem, language: ReportLanguage) -> str:
    if item.decision == "confirmed":
        return _label("confirmed", language)
    if item.decision == "dismissed":
        return _label("dismissed", language)
    return _label("unresolved", language)


def _status_fill_hex(item: ReviewItem) -> str:
    if item.decision == "confirmed":
        return _CONFIRMED_HEX
    if item.decision == "dismissed":
        return _DISMISSED_HEX
    if item.kind in ("amount_mismatch", "sum_match"):
        return _MISMATCH_HEX
    return _UNRESOLVED_HEX


def _fmt_amount(value: Optional[Decimal]) -> str:
    if value is None:
        return ""
    return f"{value:,.2f}"


def _fmt_diff(a: Optional[Decimal], b: Optional[Decimal]) -> str:
    if a is None or b is None:
        return ""
    return f"{(a - b):,.2f}"


# ---------------------------------------------------------------------------
# Summary / control totals — shared across all three formats
# ---------------------------------------------------------------------------


def _collect_side_totals(result: ReconciliationResult) -> Tuple[Decimal, Decimal, Decimal, Decimal]:
    """Returns (matched_amount_a, matched_amount_b, grand_total_a,
    grand_total_b). Grand totals are built from a row_id-keyed map, not a
    running sum, so an "ambiguous" review item's candidates — which can
    legitimately be listed under more than one row_a without being claimed
    by any of them (see matching.py's referenced_b_ids) — aren't double
    counted."""
    zero = Decimal("0")
    matched_amount_a = sum((p.row_a.amount or zero for p in result.matched), zero)
    matched_amount_b = sum((p.row_b.amount or zero for p in result.matched), zero)

    rows_a: Dict[str, Decimal] = {p.row_a.row_id: (p.row_a.amount or zero) for p in result.matched}
    rows_b: Dict[str, Decimal] = {p.row_b.row_id: (p.row_b.amount or zero) for p in result.matched}
    for item in result.reviewed:
        if item.row_a is not None:
            rows_a[item.row_a.row_id] = item.row_a.amount or zero
        if item.row_b is not None:
            rows_b[item.row_b.row_id] = item.row_b.amount or zero
        for c in item.candidates:
            rows_b[c.row_id] = c.amount or zero
        # group_match's group rows can be on either side depending on
        # direction (several of ours -> one of theirs, or the reverse) —
        # route each by its own source rather than assuming a side.
        for g in item.group:
            (rows_a if g.source == "sap_odoo" else rows_b)[g.row_id] = g.amount or zero

    grand_total_a = sum(rows_a.values(), zero)
    grand_total_b = sum(rows_b.values(), zero)
    return matched_amount_a, matched_amount_b, grand_total_a, grand_total_b


def _compute_summary(result: ReconciliationResult) -> dict:
    matched_amount_a, matched_amount_b, grand_total_a, grand_total_b = _collect_side_totals(result)

    mismatch_count = sum(1 for i in result.reviewed if i.kind == "amount_mismatch")
    group_match_count = sum(1 for i in result.reviewed if i.kind == "group_match")
    sum_match_count = sum(1 for i in result.reviewed if i.kind == "sum_match")
    other_exception_count = len(result.reviewed) - mismatch_count - group_match_count - sum_match_count
    confirmed = sum(1 for i in result.reviewed if i.decision == "confirmed")
    dismissed = sum(1 for i in result.reviewed if i.decision == "dismissed")
    unresolved = len(result.reviewed) - confirmed - dismissed

    pct_a = (matched_amount_a / grand_total_a * 100) if grand_total_a else None
    pct_b = (matched_amount_b / grand_total_b * 100) if grand_total_b else None

    return {
        "matched_count": len(result.matched),
        "matched_amount_a": matched_amount_a,
        "matched_amount_b": matched_amount_b,
        "grand_total_a": grand_total_a,
        "grand_total_b": grand_total_b,
        "pct_a": pct_a,
        "pct_b": pct_b,
        "mismatch_count": mismatch_count,
        "group_match_count": group_match_count,
        "sum_match_count": sum_match_count,
        "other_exception_count": other_exception_count,
        "confirmed": confirmed,
        "dismissed": dismissed,
        "unresolved": unresolved,
    }


def _group_summary(rows: List, language: ReportLanguage) -> Tuple[str, Decimal]:
    """Combines a group_match's several-item side into one "ref: desc; ref:
    desc" display string plus a summed amount, for the exception table's
    single-column-per-side layout."""
    zero = Decimal("0")
    parts = []
    for r in rows:
        ref = r.reference or "-"
        desc = _maybe_translate(r.description, language)
        parts.append(f"{ref}: {desc}" if desc else ref)
    total = sum((r.amount or zero for r in rows), zero)
    return "; ".join(parts), total


def _candidates_summary(rows: List, language: ReportLanguage) -> Tuple[str, Optional[Decimal]]:
    """Combines an ambiguous item's row_b candidates into one "ref: desc;
    ref: desc" display string, plus a representative amount for the
    Difference column. Every candidate already satisfies the same
    amount-tolerance check against row_a (see matching.py's
    _fuzzy_candidates), so any one of them is a fair representative — this
    isn't picking which candidate is "the" match, just which number to
    show in a single Difference cell."""
    parts = []
    for r in rows:
        ref = r.reference or "-"
        desc = _maybe_translate(r.description, language)
        parts.append(f"{ref}: {desc}" if desc else ref)
    representative_amount = rows[0].amount if rows else None
    return "; ".join(parts), representative_amount


def _exception_row(item: ReviewItem, language: ReportLanguage) -> List[str]:
    if item.kind == "ambiguous":
        candidates_text, representative_amount = _candidates_summary(item.candidates, language)
        # Two different reasons land here (see matching.py's Pass 3 / the
        # frontend's identical distinction): row_a genuinely has several
        # real candidates to choose between, OR (exactly one candidate
        # listed) that one candidate is ALSO wanted by a different row_a —
        # a static export can't cross-reference "which other row," but it
        # can at least say that's why this one isn't a clean fuzzy_match.
        kind_label = _review_kind_label(item.kind, language)
        if len(item.candidates) == 1:
            kind_label += " (also a candidate for another transaction)"
        return [
            kind_label, item.row_a.reference or "",
            _maybe_translate(item.row_a.description, language), _fmt_amount(item.row_a.amount),
            candidates_text, _fmt_amount(representative_amount),
            _fmt_diff(item.row_a.amount, representative_amount),
            _status_label(item, language), item.explanation or "",
        ]

    if item.kind in ("group_match", "sum_match"):
        # Identical rendering shape for both -- sum_match is the same
        # "single item vs. a summed group" structure as group_match, just
        # found without a connecting signal (see matching.py's Pass 5);
        # _review_kind_label already picks the right label text for
        # whichever kind this actually is.
        group_text, group_total = _group_summary(item.group, language)
        if item.row_b is not None:
            # several of ours (the group) -> one of theirs (row_b)
            single = item.row_b
            return [
                _review_kind_label(item.kind, language), single.reference or "",
                group_text, _fmt_amount(group_total),
                _maybe_translate(single.description, language), _fmt_amount(single.amount),
                _fmt_diff(group_total, single.amount),
                _status_label(item, language), item.explanation or "",
            ]
        # one of ours (row_a) -> several of theirs (the group)
        single = item.row_a
        return [
            _review_kind_label(item.kind, language), single.reference or "",
            _maybe_translate(single.description, language), _fmt_amount(single.amount),
            group_text, _fmt_amount(group_total),
            _fmt_diff(single.amount, group_total),
            _status_label(item, language), item.explanation or "",
        ]

    row_a = item.row_a
    row_b = item.row_b

    # Diagnostic only — possible_vat_gap doesn't change anything about
    # what data is shown, just replaces the generic "Amount mismatch"
    # label with a more specific, actionable one when the gap is close to
    # Saudi VAT's 15% (see matching.py's _is_vat_gap). The confirm/
    # dismiss/unresolved decision is still the reviewer's.
    if item.kind == "amount_mismatch" and item.possible_vat_gap and row_a and row_b:
        kind_label = (
            f"Possible VAT-inclusive/exclusive difference — "
            f"SAR {_fmt_amount(row_a.amount)} vs SAR {_fmt_amount(row_b.amount)} (≈15% gap)"
        )
    else:
        kind_label = _review_kind_label(item.kind, language)

    return [
        kind_label,
        row_a.reference if row_a else (row_b.reference if row_b else ""),
        _maybe_translate(row_a.description, language) if row_a else "",
        _fmt_amount(row_a.amount) if row_a else "",
        _maybe_translate(row_b.description, language) if row_b else "",
        _fmt_amount(row_b.amount) if row_b else "",
        _fmt_diff(row_a.amount if row_a else None, row_b.amount if row_b else None),
        _status_label(item, language),
        item.explanation or "",
    ]


_EXCEPTION_HEADER_KEYS = [
    "exception_type", "reference", "sap_odoo_side", "sap_odoo_amount",
    "second_doc_side", "second_doc_amount", "difference", "status", "explanation",
]


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------


def _generate_xlsx(result: ReconciliationResult, language: ReportLanguage) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    summary = _compute_summary(result)

    header_fill = PatternFill("solid", fgColor=_ACCENT_HEX)
    header_font = Font(bold=True, color=_ACCENT_TEXT_HEX, size=11)
    title_font = Font(bold=True, size=16, color=_ACCENT_HEX)
    subtitle_font = Font(italic=True, size=9, color="666666")
    section_font = Font(bold=True, size=12, color=_ACCENT_HEX)
    bold_font = Font(bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    rtl = language == "ar"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _label("reconciliation_report", language)[:31]
    ws.sheet_view.rightToLeft = rtl

    def set_row(row_idx: int, values: List, *, font: Optional[Font] = None, fill: Optional[PatternFill] = None,
                bordered: bool = False, number_cols: Optional[List[int]] = None) -> None:
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            if bordered:
                cell.border = border
            if number_cols and col_idx in number_cols:
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")

    r = 1
    ws.cell(row=r, column=1, value=_label("reconciliation_report", language)).font = title_font
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(_EXCEPTION_HEADER_KEYS))
    r += 1
    ws.cell(row=r, column=1, value=f"{_label('generated', language)}: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = subtitle_font
    r += 2

    # --- Summary / control totals -----------------------------------------
    ws.cell(row=r, column=1, value=_label("summary", language)).font = section_font
    r += 1
    set_row(r, ["", _label("sap_odoo_side", language), _label("second_doc_side", language)],
            font=header_font, fill=header_fill, bordered=True)
    r += 1
    set_row(r, [_label("grand_total", language), summary["grand_total_a"], summary["grand_total_b"]],
            bordered=True, number_cols=[2, 3])
    r += 1
    set_row(r, [_label("matched_amount", language), summary["matched_amount_a"], summary["matched_amount_b"]],
            font=bold_font, bordered=True, number_cols=[2, 3])
    r += 1
    pct_a = f"{summary['pct_a']:.1f}%" if summary["pct_a"] is not None else "—"
    pct_b = f"{summary['pct_b']:.1f}%" if summary["pct_b"] is not None else "—"
    set_row(r, [_label("matched_pct", language), pct_a, pct_b], bordered=True)
    r += 2

    set_row(r, [_label("total_matched", language), summary["matched_count"]])
    r += 1
    mismatch_cell_font = bold_font if summary["mismatch_count"] else None
    set_row(r, [_label("total_amount_mismatch", language), summary["mismatch_count"]], font=mismatch_cell_font)
    r += 1
    group_cell_font = bold_font if summary["group_match_count"] else None
    set_row(r, [_label("total_group_match", language), summary["group_match_count"]], font=group_cell_font)
    r += 1
    sum_cell_font = bold_font if summary["sum_match_count"] else None
    set_row(r, [_label("total_sum_match", language), summary["sum_match_count"]], font=sum_cell_font)
    r += 1
    set_row(r, [_label("total_other_exceptions", language), summary["other_exception_count"]])
    r += 1
    set_row(r, [_label("total_confirmed", language), summary["confirmed"]])
    r += 1
    set_row(r, [_label("total_dismissed", language), summary["dismissed"]])
    r += 1
    set_row(r, [_label("total_unresolved", language), summary["unresolved"]])
    r += 2

    # --- Matched items -------------------------------------------------------
    ws.cell(row=r, column=1, value=_label("matched_items", language)).font = section_font
    r += 1
    matched_headers = [
        _label("reference", language), _label("date", language), _label("description", language),
        _label("sap_odoo_amount", language), _label("second_doc_amount", language), _label("match_type", language),
    ]
    set_row(r, matched_headers, font=header_font, fill=header_fill, bordered=True)
    r += 1
    for pair in result.matched:
        reference = pair.row_a.reference or pair.row_b.reference or ""
        row_date = pair.row_a.date or pair.row_b.date
        description = _maybe_translate(pair.row_a.description or pair.row_b.description, language)
        set_row(r, [
            reference, str(row_date) if row_date else "", description,
            pair.row_a.amount, pair.row_b.amount, _match_type_label(pair.match_type, language),
        ], bordered=True, number_cols=[4, 5])
        r += 1
    if not result.matched:
        r += 1
    r += 1

    # --- Exceptions ------------------------------------------------------------
    ws.cell(row=r, column=1, value=_label("exceptions", language)).font = section_font
    r += 1
    exception_headers = [_label(k, language) for k in _EXCEPTION_HEADER_KEYS]
    set_row(r, exception_headers, font=header_font, fill=header_fill, bordered=True)
    r += 1
    for item in result.reviewed:
        fill = PatternFill("solid", fgColor=_status_fill_hex(item))
        row_values = _exception_row(item, language)
        set_row(r, row_values, fill=fill, bordered=True, number_cols=[4, 6, 7])
        r += 1
    if not result.reviewed:
        r += 1
    r += 1

    if result.narration:
        ws.cell(row=r, column=1, value=_label("narration", language)).font = section_font
        r += 1
        ws.cell(row=r, column=1, value=result.narration)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(_EXCEPTION_HEADER_KEYS))
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")

    widths = [16, 13, 26, 16, 16, 18, 14, 13, 26]
    for i, width in enumerate(widths[:len(_EXCEPTION_HEADER_KEYS)], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = None  # summary/matched/exceptions are stacked, not one continuous table — freezing a single header would be misleading

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------


def _set_cell_shading(cell, color_hex: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def _style_header_row(row, color_hex: str = _ACCENT_HEX, text_hex: str = _ACCENT_TEXT_HEX) -> None:
    from docx.shared import RGBColor

    for cell in row.cells:
        _set_cell_shading(cell, color_hex)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(text_hex)


def _generate_docx(result: ReconciliationResult, language: ReportLanguage) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    summary = _compute_summary(result)

    doc = docx.Document()
    title = doc.add_heading(_label("reconciliation_report", language), level=1)
    title.runs[0].font.color.rgb = RGBColor.from_string(_ACCENT_HEX)
    subtitle = doc.add_paragraph(f"{_label('generated', language)}: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.size = Pt(9)

    doc.add_heading(_label("summary", language), level=2)
    summary_table = doc.add_table(rows=1, cols=3)
    summary_table.style = "Table Grid"
    _style_header_row(summary_table.rows[0])
    summary_table.rows[0].cells[0].text = ""
    summary_table.rows[0].cells[1].text = _label("sap_odoo_side", language)
    summary_table.rows[0].cells[2].text = _label("second_doc_side", language)

    def add_summary_row(label_key: str, a, b, bold: bool = False) -> None:
        row = summary_table.add_row().cells
        row[0].text = _label(label_key, language)
        row[1].text = str(a)
        row[2].text = str(b)
        if bold:
            for cell in row:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True

    add_summary_row("grand_total", _fmt_amount(summary["grand_total_a"]), _fmt_amount(summary["grand_total_b"]))
    add_summary_row("matched_amount", _fmt_amount(summary["matched_amount_a"]), _fmt_amount(summary["matched_amount_b"]), bold=True)
    pct_a = f"{summary['pct_a']:.1f}%" if summary["pct_a"] is not None else "—"
    pct_b = f"{summary['pct_b']:.1f}%" if summary["pct_b"] is not None else "—"
    add_summary_row("matched_pct", pct_a, pct_b)

    doc.add_paragraph(f"{_label('total_matched', language)}: {summary['matched_count']}")
    p = doc.add_paragraph(f"{_label('total_amount_mismatch', language)}: {summary['mismatch_count']}")
    if summary["mismatch_count"]:
        p.runs[0].bold = True
        p.runs[0].font.color.rgb = RGBColor.from_string("C0392B")
    p_group = doc.add_paragraph(f"{_label('total_group_match', language)}: {summary['group_match_count']}")
    if summary["group_match_count"]:
        p_group.runs[0].bold = True
    p_sum = doc.add_paragraph(f"{_label('total_sum_match', language)}: {summary['sum_match_count']}")
    if summary["sum_match_count"]:
        p_sum.runs[0].bold = True
    doc.add_paragraph(f"{_label('total_other_exceptions', language)}: {summary['other_exception_count']}")
    doc.add_paragraph(f"{_label('total_confirmed', language)}: {summary['confirmed']}")
    doc.add_paragraph(f"{_label('total_dismissed', language)}: {summary['dismissed']}")
    doc.add_paragraph(f"{_label('total_unresolved', language)}: {summary['unresolved']}")

    doc.add_heading(_label("matched_items", language), level=2)
    matched_headers = [
        "reference", "date", "description", "sap_odoo_amount", "second_doc_amount", "match_type",
    ]
    table = doc.add_table(rows=1, cols=len(matched_headers))
    table.style = "Table Grid"
    for i, key in enumerate(matched_headers):
        table.rows[0].cells[i].text = _label(key, language)
    _style_header_row(table.rows[0])
    for pair in result.matched:
        row = table.add_row().cells
        reference = pair.row_a.reference or pair.row_b.reference or ""
        row_date = pair.row_a.date or pair.row_b.date
        row[0].text = reference
        row[1].text = str(row_date) if row_date else ""
        row[2].text = _maybe_translate(pair.row_a.description or pair.row_b.description, language)
        row[3].text = _fmt_amount(pair.row_a.amount)
        row[4].text = _fmt_amount(pair.row_b.amount)
        row[5].text = _match_type_label(pair.match_type, language)

    doc.add_heading(_label("exceptions", language), level=2)
    ex_table = doc.add_table(rows=1, cols=len(_EXCEPTION_HEADER_KEYS))
    ex_table.style = "Table Grid"
    for i, key in enumerate(_EXCEPTION_HEADER_KEYS):
        ex_table.rows[0].cells[i].text = _label(key, language)
    _style_header_row(ex_table.rows[0])
    for item in result.reviewed:
        row = ex_table.add_row().cells
        for i, value in enumerate(_exception_row(item, language)):
            row[i].text = str(value)
        fill = _status_fill_hex(item)
        for cell in row:
            _set_cell_shading(cell, fill)

    if result.narration:
        doc.add_heading(_label("narration", language), level=2)
        doc.add_paragraph(result.narration)

    if language == "ar":
        for paragraph in doc.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# pdf
# ---------------------------------------------------------------------------


def _generate_pdf(result: ReconciliationResult, language: ReportLanguage) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    summary = _compute_summary(result)
    styles = getSampleStyleSheet()
    subtitle_style = ParagraphStyle("subtitle", parent=styles["Normal"], textColor=colors.grey, fontSize=8)

    story = [
        Paragraph(_label("reconciliation_report", language), styles["Title"]),
        Paragraph(f"{_label('generated', language)}: {datetime.now().strftime('%Y-%m-%d %H:%M')}", subtitle_style),
        Spacer(1, 12),
    ]

    story.append(Paragraph(_label("summary", language), styles["Heading2"]))
    pct_a = f"{summary['pct_a']:.1f}%" if summary["pct_a"] is not None else "—"
    pct_b = f"{summary['pct_b']:.1f}%" if summary["pct_b"] is not None else "—"
    summary_data = [
        ["", _label("sap_odoo_side", language), _label("second_doc_side", language)],
        [_label("grand_total", language), _fmt_amount(summary["grand_total_a"]), _fmt_amount(summary["grand_total_b"])],
        [_label("matched_amount", language), _fmt_amount(summary["matched_amount_a"]), _fmt_amount(summary["matched_amount_b"])],
        [_label("matched_pct", language), pct_a, pct_b],
    ]
    summary_table = Table(summary_data, colWidths=[5 * cm, 5 * cm, 5 * cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_ACCENT_HEX}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        f"{_label('total_matched', language)}: {summary['matched_count']} &nbsp;&nbsp; "
        f"{_label('total_amount_mismatch', language)}: {summary['mismatch_count']} &nbsp;&nbsp; "
        f"{_label('total_group_match', language)}: {summary['group_match_count']} &nbsp;&nbsp; "
        f"{_label('total_sum_match', language)}: {summary['sum_match_count']} &nbsp;&nbsp; "
        f"{_label('total_other_exceptions', language)}: {summary['other_exception_count']} &nbsp;&nbsp; "
        f"{_label('total_confirmed', language)}: {summary['confirmed']} &nbsp;&nbsp; "
        f"{_label('total_dismissed', language)}: {summary['dismissed']} &nbsp;&nbsp; "
        f"{_label('total_unresolved', language)}: {summary['unresolved']}",
        styles["Normal"],
    ))
    story.append(Spacer(1, 14))

    story.append(Paragraph(_label("matched_items", language), styles["Heading2"]))
    matched_data = [[
        _label("reference", language), _label("date", language), _label("description", language),
        _label("sap_odoo_amount", language), _label("second_doc_amount", language), _label("match_type", language),
    ]]
    for pair in result.matched:
        row_date = pair.row_a.date or pair.row_b.date
        matched_data.append([
            pair.row_a.reference or pair.row_b.reference or "",
            str(row_date) if row_date else "",
            _maybe_translate(pair.row_a.description or pair.row_b.description, language),
            _fmt_amount(pair.row_a.amount), _fmt_amount(pair.row_b.amount),
            _match_type_label(pair.match_type, language),
        ])
    matched_table = Table(matched_data, colWidths=[2.8 * cm, 2.2 * cm, 4.5 * cm, 2.7 * cm, 2.7 * cm, 3.1 * cm], repeatRows=1)
    matched_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_ACCENT_HEX}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (3, 1), (4, -1), "RIGHT"),
    ]))
    story.append(matched_table)
    story.append(Spacer(1, 14))

    story.append(Paragraph(_label("exceptions", language), styles["Heading2"]))
    exception_data = [[_label(k, language) for k in _EXCEPTION_HEADER_KEYS]]
    row_colors = [colors.white]  # header placeholder, overwritten by BACKGROUND command below
    for item in result.reviewed:
        exception_data.append(_exception_row(item, language))
        row_colors.append(colors.HexColor(f"#{_status_fill_hex(item)}"))
    exception_table = Table(
        exception_data,
        colWidths=[2.4 * cm, 2 * cm, 2.6 * cm, 2 * cm, 2.6 * cm, 2 * cm, 1.8 * cm, 1.8 * cm, 3.5 * cm],
        repeatRows=1,
    )
    exception_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_ACCENT_HEX}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("ALIGN", (3, 1), (6, -1), "RIGHT"),
    ]
    for i, color in enumerate(row_colors[1:], start=1):
        exception_style.append(("BACKGROUND", (0, i), (-1, i), color))
    exception_table.setStyle(TableStyle(exception_style))
    story.append(exception_table)

    if result.narration:
        story.append(Spacer(1, 14))
        story.append(Paragraph(_label("narration", language), styles["Heading2"]))
        story.append(Paragraph(result.narration, styles["Normal"]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    doc.build(story)
    return buf.getvalue()


def generate_reconciliation_report(result: ReconciliationResult, format: ReportFormat = "xlsx", language: ReportLanguage = "en") -> bytes:
    if format == "xlsx":
        return _generate_xlsx(result, language)
    if format == "docx":
        return _generate_docx(result, language)
    if format == "pdf":
        return _generate_pdf(result, language)
    raise ValueError(f"Unsupported report format: {format!r}")
